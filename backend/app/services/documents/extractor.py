import io
import os
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import BinaryIO

import pypdf
import docx


class DocumentExtractionError(Exception):
    """Raised when document content extraction fails."""
    pass


class DocumentValidationError(Exception):
    """Raised when file validation fails."""
    pass


@dataclass
class ExtractedSection:
    page_number: int | None
    text: str


ALLOWED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"}

MIME_TYPE_MAP = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
}


def validate_file_content(filename: str, content: bytes, max_size_bytes: int | None = None) -> str:
    """Validate document extension, size, and magic bytes/integrity.
    
    Returns the resolved canonical MIME type.
    Raises DocumentValidationError if invalid.
    """
    if not filename or "." not in filename:
        raise DocumentValidationError("File has no extension")

    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentValidationError(
            f"Unsupported file extension '{ext}'. Allowed extensions: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    if max_size_bytes is not None and len(content) > max_size_bytes:
        raise DocumentValidationError(
            f"File size ({len(content)} bytes) exceeds the maximum allowed limit of {max_size_bytes} bytes."
        )

    if len(content) == 0:
        raise DocumentValidationError("File is empty (0 bytes).")

    # Content integrity check
    if ext == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise DocumentValidationError("Invalid PDF: Missing %PDF- file header signature.")
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            if len(reader.pages) == 0:
                raise DocumentValidationError("PDF document has no pages.")
        except Exception as exc:
            raise DocumentValidationError(f"Corrupted or invalid PDF file: {exc}") from exc

    elif ext == ".docx":
        if not content.startswith(b"PK\x03\x04"):
            raise DocumentValidationError("Invalid DOCX: Missing standard ZIP file header signature.")
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                namelist = zf.namelist()
                if not any("word/document.xml" in name for name in namelist):
                    raise DocumentValidationError("Invalid DOCX: File does not contain standard Word document structure.")
        except zipfile.BadZipFile as exc:
            raise DocumentValidationError(f"Corrupted DOCX archive: {exc}") from exc

    elif ext in (".txt", ".md", ".markdown"):
        # Check for null bytes indicative of binary files
        sample = content[:4096]
        if b"\x00" in sample:
            raise DocumentValidationError("Binary file detected; text and markdown files must contain text data.")
        try:
            content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                content.decode("latin-1")
            except Exception as exc:
                raise DocumentValidationError(f"Text encoding could not be decoded safely: {exc}") from exc

    return MIME_TYPE_MAP.get(ext, "application/octet-stream")


def extract_document_text(file_path: Path | str, mime_type: str | None = None) -> list[ExtractedSection]:
    """Extract sections with text and page numbers from a document file on disk.
    
    Supports PDF (page-by-page), DOCX (paragraphs and tables), and TXT/Markdown.
    """
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"File not found: {path}")

    ext = path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(path)
    elif ext == ".docx":
        return _extract_docx(path)
    elif ext in (".txt", ".md", ".markdown"):
        return _extract_text(path)
    else:
        raise DocumentExtractionError(f"Unsupported file format for extraction: {ext}")


def _extract_pdf(path: Path) -> list[ExtractedSection]:
    try:
        reader = pypdf.PdfReader(str(path))
        sections: list[ExtractedSection] = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            page_text = page_text.replace("\r\n", "\n").replace("\r", "\n").strip()
            if page_text:
                sections.append(ExtractedSection(page_number=idx + 1, text=page_text))
        if not sections:
            # If all pages were empty/scanned images
            sections.append(ExtractedSection(page_number=1, text=""))
        return sections
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract PDF text from {path.name}: {exc}") from exc


def _extract_docx(path: Path) -> list[ExtractedSection]:
    try:
        doc = docx.Document(str(path))
        paragraphs: list[str] = []

        # Extract text from paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        # Extract text from tables
        for table in doc.tables:
            table_rows: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                paragraphs.append("\n".join(table_rows))

        combined_text = "\n\n".join(paragraphs).replace("\r\n", "\n").replace("\r", "\n").strip()
        return [ExtractedSection(page_number=1, text=combined_text)]
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract DOCX text from {path.name}: {exc}") from exc


def _extract_text(path: Path) -> list[ExtractedSection]:
    try:
        raw_bytes = path.read_bytes()
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1", errors="replace")
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        return [ExtractedSection(page_number=1, text=text.strip())]
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to read text file {path.name}: {exc}") from exc

